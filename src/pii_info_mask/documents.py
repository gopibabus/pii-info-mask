"""Mask PII inside TXT, DOCX and PDF files while keeping each format intact."""

from dataclasses import dataclass, field
from pathlib import Path

import docx
import pymupdf
from docx.text.hyperlink import Hyperlink

from .detect import Span, detect, mask_text

SUPPORTED = (".txt", ".docx", ".pdf")


@dataclass
class Result:
    text: str  # extracted original text, for the preview
    spans: list[Span]
    warnings: list[str] = field(default_factory=list)


def mask_file(src: Path, dst: Path) -> Result:
    suffix = src.suffix.lower()
    if suffix == ".txt":
        return _mask_txt(src, dst)
    if suffix == ".docx":
        return _mask_docx(src, dst)
    if suffix == ".pdf":
        return _mask_pdf(src, dst)
    raise ValueError(f"Unsupported file type {suffix!r}; use one of {', '.join(SUPPORTED)}")


def _mask_txt(src: Path, dst: Path) -> Result:
    text = src.read_text(encoding="utf-8", errors="replace")
    spans = detect(text)
    dst.write_text(mask_text(text, spans), encoding="utf-8")
    return Result(text, spans)


# --- DOCX ---------------------------------------------------------------------------


def _paragraphs(document):
    """Every paragraph in body, tables (including nested) and headers/footers."""

    def walk(container):
        for para in container.paragraphs:
            yield para
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from walk(cell)

    yield from walk(document)
    for section in document.sections:
        for part in (section.header, section.footer):
            if not part.is_linked_to_previous:
                yield from walk(part)


def _runs(para):
    runs = []
    for item in para.iter_inner_content():
        runs.extend(item.runs if isinstance(item, Hyperlink) else [item])
    return runs


def _replace_in_runs(runs, start: int, end: int, replacement: str) -> None:
    """Replace chars [start, end) of the runs' joined text, keeping each run's formatting."""
    pos, placed = 0, False
    for run in runs:
        text = run.text
        run_start, run_end = pos, pos + len(text)
        pos = run_end
        if run_end <= start or run_start >= end:
            continue
        cut_from, cut_to = max(start, run_start) - run_start, min(end, run_end) - run_start
        run.text = text[:cut_from] + ("" if placed else replacement) + text[cut_to:]
        placed = True


def _mask_docx(src: Path, dst: Path) -> Result:
    document = docx.Document(str(src))
    paras = [p for p in _paragraphs(document)]
    texts = [p.text for p in paras]
    full = "\n".join(texts)
    spans = detect(full)

    offset = 0
    for para, text in zip(paras, texts):
        para_end = offset + len(text)
        local = [
            (max(s.start, offset) - offset, min(s.end, para_end) - offset, s.label)
            for s in spans
            if s.start < para_end and s.end > offset
        ]
        if local:
            runs = _runs(para)
            if "".join(r.text for r in runs) != text:
                # Text lives outside plain runs (fields, content controls): rewrite as one run.
                for run in runs[1:]:
                    run.text = ""
                runs = runs[:1] or [para.add_run()]
                runs[0].text = text
            for start, end, label in reversed(local):
                _replace_in_runs(runs, start, end, f"[{label}]")
        offset = para_end + 1  # the "\n" joiner

    document.save(str(dst))
    return Result(full, spans)


# --- PDF ----------------------------------------------------------------------------


def _mask_pdf(src: Path, dst: Path) -> Result:
    pdf = pymupdf.open(str(src))
    all_text, all_spans, warnings = [], [], []
    offset = 0

    for page in pdf:
        text = page.get_text()
        if not text.strip() and page.get_images():
            warnings.append(f"Page {page.number + 1} looks scanned (image only); it was not masked.")
        spans = detect(text)
        # A span can wrap lines and search_for matches within a line, so look up each piece.
        # search_for returns every occurrence, so each distinct piece is searched only once.
        pieces = {}
        for span in spans:
            for piece in text[span.start : span.end].splitlines():
                if piece.strip():
                    pieces.setdefault(piece.strip(), f"[{span.label}]")
        textpage = page.get_textpage()
        for piece, label in pieces.items():
            for rect in page.search_for(piece, textpage=textpage):
                page.add_redact_annot(
                    rect,
                    text=label,
                    fontsize=max(4, min(rect.height * 0.7, 10)),
                    fill=(0, 0, 0),
                    text_color=(1, 1, 1),
                )
        # Removes the underlying text, not just draws over it; leaves images untouched.
        page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE)

        all_text.append(text)
        all_spans += [Span(s.start + offset, s.end + offset, s.label) for s in spans]
        offset += len(text)

    pdf.save(str(dst), garbage=4, deflate=True)
    return Result("".join(all_text), all_spans, warnings)
